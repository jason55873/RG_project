import openpyxl
import PySimpleGUI as sg
import xlrd
import pandas as pd
from openpyxl.styles import PatternFill


# 轉換訂單
def change_excel(order_file_path, no_barcode_quotation_file_path, barcode_quotation_file_path, new_order_file_path):
    try:
        no_barcode_df = pd.read_excel(
            no_barcode_quotation_file_path, skiprows=range(3))
        no_barcode_df.iloc[:, 0] = no_barcode_df.iloc[:, 0].str.strip()

        barcode_df = pd.read_excel(
            barcode_quotation_file_path, skiprows=range(3))
        barcode_df.iloc[:, 0] = barcode_df.iloc[:, 0].str.strip()

        # 將第一個col設為字典的鍵，其他欄位設為字典的值
        data_dict = {}
        for index, row in no_barcode_df.iterrows():
            key = row.iloc[0]  # 第一個欄位作為鍵
            value = row[1:].to_dict()  # 其他欄位轉為字典作為值
            data_dict[key] = value

        for index, row in barcode_df.iterrows():
            key = row.iloc[0]  # 第一個欄位作為鍵
            value = row[1:].to_dict()  # 其他欄位轉為字典作為值
            data_dict[key] = value

        # 使用 Pandas 读取 .xls 文件
        df = pd.read_excel(order_file_path, skiprows=2)
        df['商品代號'] = df['商品代號'].str.strip()
        df['產品編號'] = df['商品代號'].apply(lambda x: str(
            data_dict.get(x, {}).get('Unnamed: 1', '')))
        df['價格'] = df['商品代號'].apply(lambda x: str(
            data_dict.get(x, {}).get('進價', '')))

        grouped = df.groupby('現場名稱', group_keys=False)

        def custom_sort_key(s):
            # 移除多余的字符
            s = s.split('、')[0].split('-')[0].split('.')[0]
            # 首先按照字符串长度排序，然后按照字典顺序排序
            return (len(s), s)

        def custom_order(mark):
            mark = sorted(mark, key=custom_sort_key)
            return mark

        # 访问每个分组
        with pd.ExcelWriter(new_order_file_path) as writer:
            totle_bill = 0
            for name, group in grouped:
                total_quantity = 0
                total_amount = 0
                custom_order_list = custom_order(group['產品編號'])
                group['產品編號'] = pd.Categorical(
                    group['產品編號'], categories=custom_order_list, ordered=True)
                sorted_group = group.sort_values(by='產品編號')
                total_quantity = sorted_group["數量"].sum()
                total_amount = (pd.to_numeric(
                    sorted_group['數量']) * pd.to_numeric(sorted_group["價格"])).sum()
                totle_bill += total_amount
                new_data = {'數量': total_quantity, '價格': total_amount}
                sorted_group = pd.concat(
                    [sorted_group, pd.DataFrame([new_data])], ignore_index=True)

                # 重新排序
                new_order_of_columns = [
                    '現場名稱', '商品代號', '條碼', '機型', '商品名稱', '產品編號', '數量', '價格', '備註']  # 新的欄位順序
                sorted_group = sorted_group.reindex(
                    columns=new_order_of_columns)

                # 将每个分组排序后的数据写入不同的工作表
                sorted_group.to_excel(writer, sheet_name=name, index=False)

            # 統計整份訂單資訊
            bill_info = pd.DataFrame(
                [{'總家數': len(grouped), '訂單總金額': totle_bill}])
            bill_info.to_excel(writer, sheet_name='訂單資訊', index=False)

        return {'status': 'success'}
    except Exception as e:
        return {'status': 'error', 'error_msg': e.__str__()}

# 標示新條碼


def find_new_tag(order_file_path, quotation_file_path, save_mark_file_path):
    try:
        # 建立字典，存儲商品編號和條碼對應的資料
        a_data_dict = {}

        if order_file_path.endswith('.xls'):
            # 讀取訂單檔案
            a_wb = xlrd.open_workbook(order_file_path)
            sheet = a_wb.sheet_by_index(0)

            num_rows = sheet.nrows
            num_cols = sheet.ncols

            data = []
            for row_index in range(4, num_rows):
                row_data = []
                product_number = sheet.cell_value(row_index, 1).replace(
                    " ", "") if sheet.cell_value(row_index, 1) else ''  # 第二colum的商品編號
                barcode = sheet.cell_value(row_index, 5)  # 第六colum的條碼
                a_data_dict[product_number] = barcode
        else:
            # 讀取訂單檔案
            a_wb = openpyxl.load_workbook(order_file_path)
            a_ws = a_wb.active

            for row in a_ws.iter_rows(min_row=4, max_col=6, values_only=True):
                product_number = row[1].replace(
                    " ", "") if row[1] else ""  # 第二colum的商品編號
                barcode = row[5]  # 第六colum的條碼
                a_data_dict[product_number] = barcode

        # 讀取報價單檔案
        b_wb = openpyxl.load_workbook(quotation_file_path)
        b_ws = b_wb.active

        # 创建一个白色填充，用于清空儲存格的背景颜色
        white_fill = PatternFill()

        # 創建一個紅色填充
        red_fill = PatternFill(start_color='FFFF0000',
                               end_color='FFFF0000',
                               fill_type='solid')
        # 更新B檔案中的資料
        for row in b_ws.iter_rows(min_row=5):
            product_number_b = row[0]  # B檔案的第一colum的編號
            # 檢查湯姆熊編號是否為空
            if product_number_b.value in a_data_dict and not row[8].value:
                barcode_b = a_data_dict[product_number_b.value]
                row[8].value = barcode_b
                for cell in row:  # 選擇整個row
                    cell.fill = red_fill  # 將背景色設為紅色
            elif row[8].value:  # 若已經有條碼則取消背景顏色
                for cell in row:  # 選擇整個row
                    cell.fill = white_fill  # 將背景色設為白色

        # 儲存更新後的B檔案
        b_wb.save(save_mark_file_path)

        return {'status': 'success'}
    except Exception as e:
        return {'status': 'error', 'error_msg': e.__str__()}


def main():
    # 設定視窗樣式
    sg.ChangeLookAndFeel('LightBlue')

    # 回傳text欄位
    def set_text(text='', key=None, size=18):
        return sg.Text(text=text, key=key, size=size, font=(12, 12))

    # 回傳input欄位
    def set_input(key=None):
        return sg.In(enable_events=True, key=key, size=(50, 1), font=(12, 12), readonly=True, disabled_readonly_background_color='lightgrey')

    # 回傳filesbrowse欄位
    def set_filesbrowse():
        return sg.FilesBrowse('選擇檔案路徑', font=(12, 12))

    # 回傳saveas欄位
    def set_saveas():
        return sg.SaveAs('選擇存擋路徑', font=(12, 12), default_extension='.xlsx')

    # 回傳frame欄位
    def set_frame(title, layout):
        return sg.Frame(title=title, layout=layout, title_color='darkblue', font=(18, 12))

    # 回傳button欄位
    def set_button(button_text, key):
        return sg.Button(button_text=button_text, key=key, font=(12, 12))

    # 設定視窗欄位
    input_layout = [
        [set_text('新品報價單路徑：'), set_input(
            "-quotation_file-"), set_filesbrowse()],
        [set_text('(無條碼)報價單路徑：'), set_input(
            "-no_barcode_quotation_file-"), set_filesbrowse()],
        [set_text('(有條碼)報價單路徑：'), set_input(
            "-barcode_quotation_file-"), set_filesbrowse()],
    ]

    output_layout = [
        [set_text('標示新條碼路徑：'), set_input("-save_mark_file-"), set_saveas()],
        [set_text('已轉換訂單路徑：'), set_input("-new_order_file-"), set_saveas()]
    ]

    result_layout = [
        [sg.Multiline(key='-status-', disabled=True,
                      background_color='Lightblue', size=(110, 5))]
    ]

    order_lay_out = [[set_text('訂單路徑：'), set_input(
        "-order_file-"), set_filesbrowse()]]

    layout = [
        [set_frame('訂單路徑', order_lay_out)],
        [set_frame('報價單檔案路徑', input_layout)],
        [set_frame('輸出檔案路徑', output_layout)],
        [set_frame('執行結果', result_layout)],
        [set_button('執行', '-start-'), set_button('關閉程式', '-EXIT-')]
    ]

    # 產生視窗
    window = sg.Window('訂單轉換', layout, finalize=True,
                       size=(750, 500), grab_anywhere=True)

    all_key_list = ["-order_file-", "-quotation_file-", "-no_barcode_quotation_file-",
                    "-barcode_quotation_file-", "-save_mark_file-", "-new_order_file-"]

    while True:
        event, values = window.read()

        if event == sg.WIN_CLOSED or event == '-EXIT-':
            break

        if event == '-start-':
            order_file_path = values["-order_file-"]
            quotation_file_path = values["-quotation_file-"]
            no_barcode_quotation_file_path = values["-no_barcode_quotation_file-"]
            barcode_quotation_file_path = values["-barcode_quotation_file-"]
            save_mark_file_path = values["-save_mark_file-"]
            new_order_file_path = values["-new_order_file-"]

            if order_file_path and quotation_file_path and no_barcode_quotation_file_path and barcode_quotation_file_path and save_mark_file_path and new_order_file_path:
                output = ''
                window['-order_file-'].update('')
                window['-quotation_file-'].update('')
                # 執行新條碼標示
                result_mark = find_new_tag(
                    order_file_path, quotation_file_path, save_mark_file_path)
                # 執行訂單轉換
                result_order = change_excel(
                    order_file_path, no_barcode_quotation_file_path, barcode_quotation_file_path, new_order_file_path)

                if result_mark['status'] == 'success' and result_order['status'] == 'success':
                    output = '執行完成，請再次選擇檔案或關閉程式'
                if result_mark['status'] != 'success':
                    output += f'新條碼標示失敗, 錯誤訊息：{result_mark["error_msg"]}\n'
                if result_order['status'] != 'success':
                    output += f'訂單轉換失敗, 錯誤訊息：{result_order["error_msg"]}\n'

                # 清空所有欄位值
                [window.find_element(k).Update('') for k in all_key_list]
                window['-status-'].Update(output)

            else:
                output = '請先選擇檔案'
                window['-status-'].update(output, font=16)

    window.close()


if __name__ == "__main__":
    main()
